# -*- coding: utf-8 -*-
"""校验 --bulk-normal 生成的批量行（create / account）。

只读 Excel，不发送。检查点：
  * 用例数量（normal 扩到 N，error/destroy 不变）
  * 账号唯一、序号连续、UniqueAccount 跟随 FAccount
  * create 的 Ref 逐行唯一；account 的 Pwd 每行能解回明文
  * 报文结构与真实流量一致
"""
import os
import re
import sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "interfaces"))

import excel_loader as XL
import create as C
import account as A
from pwd_encode import decode_pwd

OK = FAIL = 0
N = int(os.environ.get("BULK_N", "10000"))


def check(cond, label, extra=""):
    global OK, FAIL
    if cond:
        OK += 1
        print("  [OK]   %s%s" % (label, ("  " + extra) if extra else ""))
    else:
        FAIL += 1
        print("  [FAIL] %s%s" % (label, ("  " + extra) if extra else ""))


def read_excel(path):
    """读 Excel 成 [{列:值}]（单行表头 "中文\\n(key)"）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [XL._key(h) for h in next(it)]
    rows = [dict(zip(header, r)) for r in it
            if r and any(v is not None for v in r)]
    wb.close()
    return rows


def check_accounts(bulk, prefix, label):
    faccts = [str(r.get("Account_FAccount")) for r in bulk]
    check(len(set(faccts)) == len(bulk), "%s FAccount 全部唯一" % label,
          "%d 个唯一" % len(set(faccts)))
    seqs = sorted(int(f[6:]) for f in faccts)
    check(seqs == list(range(seqs[0], seqs[0] + len(bulk))),
          "%s 账号序号连续无空洞" % label, "%d~%d" % (seqs[0], seqs[-1]))
    bad = [r for r in bulk
           if str(r.get("UniqueAccount")) !=
           "%s_%s_%s" % (r.get("Account_FAccount"),
                         r.get("Account_AccountType"), r.get("Account_AccAtt"))]
    check(not bad, "%s UniqueAccount 每行跟随 FAccount" % label,
          "错 %d 行" % len(bad))
    return faccts


print("=" * 74)
print("A) create.xlsx")
print("=" * 74)
path_c = XL.default_excel("create")
print("   %s  (%.1f MB)" % (path_c, os.path.getsize(path_c) / 1e6))
pool_c = XL.load_cases(path_c, quiet=True)
cnt = Counter(t for _, t, _, _, _ in pool_c)
print("   ", dict(cnt))
check(cnt.get("normal", 0) == N, "normal = %d" % N, str(cnt.get("normal")))
check(cnt.get("error", 0) == 9, "error = 9（原有）")
check(cnt.get("destroy", 0) == 287, "destroy = 287（原有）")

rows_c = read_excel(path_c)
bulk_c = [r for r in rows_c if str(r.get("case_no", "")).startswith("CB")]
check(len(bulk_c) == N, "定位到 %d 条批量行" % N, str(len(bulk_c)))
check_accounts(bulk_c, "CB", "create")

refs = [str(r.get("Ref")) for r in bulk_c]
check(len(set(refs)) == N, "Ref 全部唯一", "%d 个唯一" % len(set(refs)))
m = re.compile(r"__REF(\d+)__$")
nums = [int(m.match(x).group(1)) for x in refs if m.match(x)]
check(len(nums) == N and sorted(nums) == list(range(1, N + 1)),
      "Ref 是 __REF1__~__REF%d__" % N)

print()
print("B) account.xlsx")
print("=" * 74)
path_a = XL.default_excel("account")
print("   %s  (%.1f MB)" % (path_a, os.path.getsize(path_a) / 1e6))
pool_a = XL.load_cases(path_a, quiet=True)
cnt_a = Counter(t for _, t, _, _, _ in pool_a)
print("   ", dict(cnt_a))
check(cnt_a.get("normal", 0) == N, "normal = %d" % N, str(cnt_a.get("normal")))
check(cnt_a.get("error", 0) == 7, "error = 7（原有）")
check(cnt_a.get("destroy", 0) == 96, "destroy = 96（原有）")

rows_a = read_excel(path_a)
bulk_a = [r for r in rows_a if str(r.get("case_no", "")).startswith("AB")]
check(len(bulk_a) == N, "定位到 %d 条批量行" % N, str(len(bulk_a)))
check_accounts(bulk_a, "AB", "account")

# 账号号段两边一致（create 与 account 应对得上）
seqs_c = sorted(int(str(r.get("Account_FAccount"))[6:]) for r in bulk_c)
seqs_a = sorted(int(str(r.get("Account_FAccount"))[6:]) for r in bulk_a)
check(seqs_c == seqs_a, "create 与 account 的账号号段完全一致",
      "%d~%d" % (seqs_a[0], seqs_a[-1]))

print()
print("=" * 74)
print("C) account 报文：Pwd 每行能解回明文 + 结构核对")
print("=" * 74)
import json
sample = pool_a[:200]
bad_pwd = bad_struct = 0
for no, typ, desc, txt, _rn in sample:
    try:
        o = json.loads(txt)
    except Exception:
        bad_struct += 1
        continue
    a = o.get("Account")
    if not isinstance(a, dict):
        bad_struct += 1
        continue
    for k in ("Account", "ClientName", "Pwd", "TradeAccount",
              "UniqueAccount", "BranchNo", "Shareholders"):
        if k not in a:
            bad_struct += 1
            break
    facct = (a.get("Account") or {}).get("FAccount")
    try:
        if decode_pwd(a.get("Pwd"), facct) != A.REAL_PWD_PLAIN:
            bad_pwd += 1
    except Exception:
        bad_pwd += 1
check(bad_struct == 0, "抽 200 条结构完整", "异常 %d" % bad_struct)
check(bad_pwd == 0, "抽 200 条 Pwd 都能解回 %r" % A.REAL_PWD_PLAIN,
      "错 %d" % bad_pwd)

sh_lens = Counter()
for no, typ, desc, txt, _rn in sample:
    try:
        a = json.loads(txt)["Account"]
    except Exception:
        continue
    for s in (a.get("Shareholders") or []):
        sa = str(s.get("SAccount", ""))
        sh_lens["A+%d位" % (len(sa) - 1) if sa.startswith("A")
                else "纯数字%d位" % len(sa)] += 1
check(set(sh_lens) == {"A+9位", "纯数字10位"},
      "股东号格式对齐真实样本（A+9位 / 10位数字）", str(dict(sh_lens)))

print()
print("=" * 74)
print("D) 抽样看报文")
print("=" * 74)
for no, typ, desc, txt, _rn in pool_a[:2]:
    a = json.loads(txt)["Account"]
    print("   [%s] %s" % (no, desc))
    print("      FAccount=%s UniqueAccount=%s" % (a["Account"].get("FAccount"),
                                                 a.get("UniqueAccount")))
    print("      TradeAccount=%s ClientName=%s" % (a.get("TradeAccount"),
                                                  a.get("ClientName")))
    print("      Pwd=%s..." % str(a.get("Pwd"))[:32])
    print("      Shareholders=%s" % a.get("Shareholders"))

print()
print("=" * 74)
print("E) modify / remove / pwdUpdate（1 万行）")
print("=" * 74)

SPECS = [
    # (接口, 批量行前缀, error数, destroy数, 单号列名)
    ("modify", "MB", 5, 82, "Ref"),
    ("remove", "RB", 5, 34, "Ref"),
    ("pwdUpdate", "PB", 5, 41, None),
]
bulk_by_iface = {}
for iface, pfx, n_err, n_dst, refcol in SPECS:
    print()
    print("--- %s ---" % iface)
    path = XL.default_excel(iface)
    pool = XL.load_cases(path, quiet=True)
    c = Counter(t for _, t, _, _, _ in pool)
    print("    %s  (%.1f MB)  %s" % (path, os.path.getsize(path) / 1e6, dict(c)))
    check(c.get("normal", 0) == N, "%s normal = %d" % (iface, N), str(c.get("normal")))
    check(c.get("error", 0) == n_err, "%s error = %d（原有）" % (iface, n_err))
    check(c.get("destroy", 0) == n_dst, "%s destroy = %d（原有）" % (iface, n_dst))

    rows_i = read_excel(path)
    bulk_i = [r for r in rows_i if str(r.get("case_no", "")).startswith(pfx)]
    check(len(bulk_i) == N, "%s 定位到 %d 条批量行" % (iface, N), str(len(bulk_i)))
    check_accounts(bulk_i, iface, iface)
    bulk_by_iface[iface] = bulk_i

print()
print("=" * 74)
print("F) ★ 单号对齐：第 i 行必须引用 create 第 i 行建出的单号")
print("=" * 74)
# create 的批量行 Ref 是 __REF1__..__REF{N}__（相对序号）
m = re.compile(r"__REF(\d+)__$")
create_refs = []
for r in bulk_c:
    mm = m.match(str(r.get("Ref")))
    create_refs.append(int(mm.group(1)) if mm else None)
check(None not in create_refs, "create 的 Ref 都能解析出序号")
check(create_refs == list(range(1, N + 1)),
      "create 的 Ref 序号正好是 1..%d（按行顺序）" % N)

for iface in ("modify", "remove"):
    refs_i = []
    for r in bulk_by_iface[iface]:
        mm = m.match(str(r.get("Ref")))
        refs_i.append(int(mm.group(1)) if mm else None)
    check(None not in refs_i, "%s 的 Ref 都能解析出序号" % iface)
    check(refs_i == create_refs,
          "★ %s 第 i 行引用的单号 == create 第 i 行建出的单号" % iface,
          "前 5: %s" % refs_i[:5] if refs_i[:5] != create_refs[:5] else "")
    # 账号也要对齐（同一个账号改/删自己建的单）
    fa_c = [str(r.get("Account_FAccount")) for r in bulk_c]
    fa_i = [str(r.get("Account_FAccount")) for r in bulk_by_iface[iface]]
    check(fa_i == fa_c, "%s 第 i 行的账号 == create 第 i 行的账号" % iface)

# pwdUpdate 无 Ref，但账号要对齐
fa_p = [str(r.get("Account_FAccount")) for r in bulk_by_iface["pwdUpdate"]]
check(fa_p == [str(r.get("Account_FAccount")) for r in bulk_c],
      "pwdUpdate 账号号段与 create 对齐")

print()
print("=" * 74)
print("G) 报文形态核对")
print("=" * 74)
for iface, refcol in (("modify", "Ref"), ("remove", "Ref"), ("pwdUpdate", None)):
    pool_i = XL.load_cases(XL.default_excel(iface), want_types={"normal"}, quiet=True)
    bad = 0
    for no, typ, desc, txt, _rn in pool_i[:200]:
        try:
            o = json.loads(txt)
        except Exception:
            bad += 1
            continue
        body = o.get(iface)
        if not isinstance(body, dict):
            bad += 1
            continue
        need = ["Account", "UniqueAccount"] + ([refcol] if refcol else [])
        for k in need:
            if k not in body:
                bad += 1
                break
    check(bad == 0, "%s 抽 200 条结构完整" % iface, "异常 %d" % bad)

# pwdUpdate 每行密文能解回
pool_p = XL.load_cases(XL.default_excel("pwdUpdate"),
                       want_types={"normal"}, quiet=True)
bad_pwd = 0
for no, typ, desc, txt, _rn in pool_p[:200]:
    try:
        a = json.loads(txt)["pwdUpdate"]
        facct = (a.get("Account") or {}).get("FAccount")
        if decode_pwd(a.get("Pwd"), facct) != A.REAL_PWD_PLAIN:
            bad_pwd += 1
    except Exception:
        bad_pwd += 1
check(bad_pwd == 0, "pwdUpdate 抽 200 条 Pwd 能解回 %r" % A.REAL_PWD_PLAIN,
      "错 %d" % bad_pwd)

# Ref 文本确实不同（避免又变成循环复用同一条）
for iface in ("modify", "remove"):
    pool_i = XL.load_cases(XL.default_excel(iface),
                           want_types={"normal"}, quiet=True)
    refs = []
    for no, typ, desc, txt, _rn in pool_i:
        try:
            refs.append(json.loads(txt)[iface].get("Ref"))
        except Exception:
            pass
    check(len(set(refs)) == N, "%s 的 %d 条 Ref 各不相同" % (iface, N),
          "%d 个唯一" % len(set(refs)))

print()
print("=" * 74)
print("H) 抽样看报文")
print("=" * 74)
for iface in ("modify", "remove", "pwdUpdate"):
    pool_i = XL.load_cases(XL.default_excel(iface),
                           want_types={"normal"}, quiet=True)
    no, typ, desc, txt, _rn = pool_i[0]
    print("   [%s] %s" % (no, desc))
    print("      %s" % txt[:170])

print()
print("=" * 74)
print("结果: %s   通过 %d / 失败 %d" % ("全部通过" if FAIL == 0 else "有失败", OK, FAIL))
print("=" * 74)
sys.exit(1 if FAIL else 0)
